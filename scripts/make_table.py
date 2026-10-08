#!/usr/bin/env python3
"""corpus.json -> the reading/annotation table (CSV, or xlsx when openpyxl is present).

    python make_table.py lit/corpus.json -o lit/corpus.csv
    python make_table.py lit/corpus.json -o lit/corpus.xlsx --xlsx

Retrieval fills the left-hand columns. The four annotation columns - threat_model, mechanism,
evaluation, limitations - come out empty unless a record already carries them, because they are
judgements made while reading, not things an API returns. Emitting them empty is the point: the
sheet is the worksheet.

CSV path is stdlib only. --xlsx needs openpyxl and falls back to CSV with a warning if it is missing.
"""

import argparse
import csv
import json
import pathlib
import sys

COLUMNS = [
    # provenance and identity
    "key", "title", "authors", "year",
    "venue", "venue_short", "tier", "peer_reviewed",
    # retrieval judgement
    "track", "facets", "relevance", "overlap_user_work",
    # annotation - filled while reading
    "threat_model", "mechanism", "evaluation", "limitations",
    # evidence
    "arxiv_id", "doi", "verified", "verified_by", "url", "source_queries", "note",
]

LIST_COLS = {"authors", "facets", "verified_by", "source_queries", "track"}


def flat(rec, col):
    v = rec.get(col)
    if col == "track":
        v = rec.get("tracks") or rec.get("track")
    if col == "facets":
        v = rec.get("facets") or rec.get("facet")
    if v is None:
        return ""
    if isinstance(v, (list, tuple, set)):
        return "; ".join(str(x) for x in v)
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    return str(v)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("corpus")
    ap.add_argument("-o", "--out", default="corpus.csv")
    ap.add_argument("--xlsx", action="store_true", help="write a styled .xlsx (needs openpyxl)")
    args = ap.parse_args()

    recs = json.loads(pathlib.Path(args.corpus).read_text(encoding="utf-8"))
    if isinstance(recs, dict):
        recs = recs.get("records") or recs.get("corpus") or []
    rows = [[flat(r, c) for c in COLUMNS] for r in recs]
    out = pathlib.Path(args.out)

    if args.xlsx:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            print("openpyxl not available - writing CSV instead", file=sys.stderr)
            args.xlsx = False
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "corpus"
            ws.append(COLUMNS)
            for row in rows:
                ws.append(row)

            head = Font(bold=True, color="FFFFFF")
            fill_prov = PatternFill("solid", fgColor="31506E")
            fill_note = PatternFill("solid", fgColor="7A4E2D")   # annotation columns stand out
            for i, c in enumerate(COLUMNS, 1):
                cell = ws.cell(row=1, column=i)
                cell.font = head
                cell.fill = fill_note if c in (
                    "threat_model", "mechanism", "evaluation", "limitations") else fill_prov
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                width = {"title": 58, "authors": 34, "venue": 40, "note": 46,
                         "source_queries": 40}.get(c, 15)
                ws.column_dimensions[get_column_letter(i)].width = width
            ws.freeze_panes = "C2"
            ws.auto_filter.ref = ws.dimensions
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=COLUMNS.index("title") + 1).alignment = Alignment(
                    wrap_text=True, vertical="top")
            wb.save(out)
            print("%d rows -> %s" % (len(rows), out))
            return

    if out.suffix == ".xlsx":
        out = out.with_suffix(".csv")
    # utf-8-sig so Excel on a cp950 host opens the diacritics correctly instead of mojibake.
    with open(out, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(COLUMNS)
        w.writerows(rows)
    print("%d rows -> %s" % (len(rows), out))


if __name__ == "__main__":
    main()
